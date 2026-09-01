"""Carga del Anexo Finagro (`config/catalogo_destinos.xlsx`) en memoria.

El Anexo es la hoja `Destinos` con esta estructura (a partir de la fila 6):

| Col | Contenido                                   | Mapeo sección 5         |
|-----|---------------------------------------------|-------------------------|
| B   | Actividad financiable (categoría macro)     | — (solo para acotar)    |
| C   | Destino (código numérico de 6 dígitos)      | `cod_rubro`             |
| D   | Destino (nombre)                            | `descripcion_rubro`     |
| E   | Producto relacionado                        | `producto_relacionado`  |
| F   | Plazo de ejecución (días calendario)        | `plazo_dias`            |
| G   | Línea de Crédito (Inversión / Cap. trabajo) | `linea_credito`         |

La categoría macro (col B) solo aparece en la primera fila de cada bloque, así
que se arrastra hacia abajo. Las filas de pie de página (notas, fecha de
publicación) no traen código numérico en la col C y se descartan.

El catálogo es de solo lectura una vez cargado; `code_resolver` lo consulta.
"""

from __future__ import annotations

from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from ..logging_conf import get_logger

logger = get_logger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent.parent
CATALOGO_PATH = PROJECT_ROOT / "config" / "catalogo_destinos.xlsx"

# Estructura física del Anexo. Si Finagro cambia el layout, ajustar aquí.
HOJA = "Destinos"
PRIMERA_FILA = 6  # las filas 1-5 son títulos/encabezados
# Índices 0-based dentro de la tupla de valores de cada fila.
_COL_CATEGORIA = 1  # B
_COL_COD_DESTINO = 2  # C
_COL_DESTINO = 3  # D
_COL_PRODUCTO = 4  # E
_COL_PLAZO = 5  # F
_COL_LINEA = 6  # G


def _clean(value: object) -> str | None:
    """Normaliza texto del Excel: descarta vacíos y espacios duros (\\xa0)."""
    if value is None:
        return None
    text = str(value).replace("\xa0", " ").strip()
    return text or None


class CatalogoEntry(BaseModel):
    """Una fila del Anexo Finagro: un destino de crédito y sus atributos."""

    categoria_macro: str
    cod_destino: int
    destino: str
    producto_relacionado: str | None = None
    plazo_dias: int | None = None
    linea_credito: str | None = None

    def resumen(self) -> str:
        """Línea compacta para enumerar en el prompt del `code_resolver`."""
        partes = [f"{self.cod_destino} — {self.destino}"]
        if self.linea_credito:
            partes.append(f"línea: {self.linea_credito}")
        if self.producto_relacionado:
            partes.append(f"producto: {self.producto_relacionado}")
        return " | ".join(partes)


class Catalogo(BaseModel):
    """Catálogo completo con índices para acotar por categoría macro."""

    entries: list[CatalogoEntry] = Field(default_factory=list)

    @property
    def categorias(self) -> list[str]:
        """Categorías macro en orden de aparición, sin duplicados."""
        vistas: list[str] = []
        for e in self.entries:
            if e.categoria_macro not in vistas:
                vistas.append(e.categoria_macro)
        return vistas

    def by_categoria(self, categoria: str) -> list[CatalogoEntry]:
        return [e for e in self.entries if e.categoria_macro == categoria]

    def by_cod_destino(self, cod: int) -> CatalogoEntry | None:
        return next((e for e in self.entries if e.cod_destino == cod), None)

    def __len__(self) -> int:
        return len(self.entries)


def load_catalogo(path: Path | None = None) -> Catalogo:
    """Carga el Anexo Finagro en un `Catalogo`.

    Arrastra la categoría macro de la col B y descarta filas sin código de
    destino numérico (encabezados intermedios y notas de pie).
    """
    path = path or CATALOGO_PATH
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[HOJA] if HOJA in wb.sheetnames else wb.active
    if ws is None:
        raise ValueError(f"El catálogo {path} no tiene ninguna hoja legible")

    entries: list[CatalogoEntry] = []
    categoria_actual: str | None = None
    descartadas = 0

    for row in ws.iter_rows(min_row=PRIMERA_FILA, values_only=True):
        categoria_celda = _clean(row[_COL_CATEGORIA])
        cod_celda = row[_COL_COD_DESTINO]

        # Una fila de datos tiene un código de destino entero en la col C.
        # Las notas de pie traen texto/fecha o None → se descartan, pero NO
        # actualizan la categoría arrastrada.
        if not isinstance(cod_celda, int):
            if categoria_celda and cod_celda is None:
                descartadas += 1
            continue

        if categoria_celda:
            categoria_actual = categoria_celda

        destino = _clean(row[_COL_DESTINO])
        if categoria_actual is None or destino is None:
            descartadas += 1
            continue

        plazo = row[_COL_PLAZO]
        entries.append(
            CatalogoEntry(
                categoria_macro=categoria_actual,
                cod_destino=int(cod_celda),
                destino=destino,
                producto_relacionado=_clean(row[_COL_PRODUCTO]),
                plazo_dias=int(plazo) if isinstance(plazo, (int, float)) else None,
                linea_credito=_clean(row[_COL_LINEA]),
            )
        )

    wb.close()
    catalogo = Catalogo(entries=entries)
    logger.info(
        "catalogo.cargado",
        destinos=len(entries),
        categorias=len(catalogo.categorias),
        descartadas=descartadas,
    )
    return catalogo


@lru_cache(maxsize=1)
def get_catalogo() -> Catalogo:
    """Catálogo cacheado — evita releer el xlsx en cada resolución."""
    return load_catalogo()
