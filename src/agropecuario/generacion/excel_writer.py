"""Renderiza la solicitud de crédito Bancolombia/Finagro.

Estrategia: abrimos el template oficial (`templates/credito_agro_template.xlsx`)
y escribimos en celdas específicas según `config/excel_cells.yaml`. Preservamos
fórmulas (=SUM), formato y celdas combinadas.
"""

from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from ..logging_conf import get_logger

logger = get_logger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent.parent
TEMPLATE_PATH = PROJECT_ROOT / "templates" / "credito_agro_template.xlsx"
CELL_MAP_PATH = PROJECT_ROOT / "config" / "excel_cells.yaml"


def render_excel(
    fields: dict[str, Any],
    output_path: Path,
    template_path: Path | None = None,
    cell_map_path: Path | None = None,
) -> Path:
    """Rellena el template con los campos provistos.

    `fields` es un dict plano con los IDs de `excel_cells.yaml` como keys.
    Para la tabla de actividades se espera `fields["actividades"] = [{...}, ...]`.
    """
    template_path = template_path or TEMPLATE_PATH
    cell_map_path = cell_map_path or CELL_MAP_PATH
    cell_map = yaml.safe_load(cell_map_path.read_text(encoding="utf-8"))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template_path, output_path)
    wb = load_workbook(str(output_path))
    ws = wb[cell_map["hoja"]]

    _write_simple_fields(ws, fields, cell_map)
    _write_id_beneficiario(ws, fields, cell_map.get("identificacion_beneficiario", {}))
    _write_text_blocks(ws, fields, cell_map)
    _write_section_4(ws, fields, cell_map)
    _write_actividades_table(ws, fields, cell_map["tabla_actividades"])
    _write_cronograma(ws, fields, cell_map)
    _write_section_6_7(ws, fields, cell_map)
    _write_digit_fields(ws, fields, cell_map.get("campos_por_digito", {}))
    # Sección 10 se deja en blanco intencionalmente.

    wb.save(str(output_path))

    # Lo que se extrajo y este template no sabe recoger. No es un fallo: la
    # extracción va por delante del papel a propósito, y esta lista es el
    # inventario de lo que habría que mapear al pasar al formulario vigente.
    huerfanos = campos_sin_celda(fields, cell_map)
    if huerfanos:
        logger.info(
            "excel.campos_sin_celda",
            campos=huerfanos,
            template=cell_map.get("template", {}).get("archivo", str(template_path)),
        )

    logger.info("excel.rendered", path=str(output_path))
    return output_path


def campos_cubiertos(cell_map: dict) -> set[str]:
    """IDs de campo que ESTE template sabe pintar, según su mapa de celdas.

    Es el espejo de lo que consumen las funciones `_write_*`: si se añade un
    bloque nuevo al writer, hay que declararlo aquí. Sirve para responder la
    pregunta que importa al cambiar de formulario: *¿qué estamos extrayendo que
    el papel todavía no recoge?*
    """
    cubiertos = {f for f in cell_map.get("campos", {}) if not f.startswith("marca_")}
    cubiertos |= set(cell_map.get("marcas") or {})
    cubiertos |= set(cell_map.get("campos_por_digito") or {})
    if cell_map.get("identificacion_beneficiario"):
        cubiertos |= {"beneficiario_id_tipo", "beneficiario_id_numero"}
    cubiertos |= {
        "forma_de_llegar",
        "justificacion_tecnica",
        "modalidad_pago_capital",
        "modalidad_pago_intereses",
        "cronograma_fecha_inicial",
        "cronograma_fecha_final",
        "actividad_economica_descripcion",
        "fag_cobertura_pct",
        "garantia_fag",
        "actividades",
    }
    return cubiertos


def campos_sin_celda(fields: dict[str, Any], cell_map: dict) -> list[str]:
    """Campos CON valor que este template no puede mostrar (orden estable)."""
    cubiertos = campos_cubiertos(cell_map)
    return sorted(
        f for f, v in fields.items()
        if f not in cubiertos and v not in (None, "", [], {})
    )


# --- helpers ---------------------------------------------------------------

def _set(ws, coord: str, value: Any) -> None:
    """Escribe en celda manejando combinadas: usa la esquina superior-izquierda.

    Si la celda solicitada está dentro de un rango combinado pero no es el origen,
    redirige al origen en lugar de fallar silenciosamente.
    """
    if value is None or value == "":
        return
    cell = ws[coord]
    for mr in ws.merged_cells.ranges:
        if coord in mr:
            cell = ws.cell(row=mr.min_row, column=mr.min_col)
            break
    cell.value = value


def _write_simple_fields(ws, fields: dict, cell_map: dict) -> None:
    campos_map = cell_map["campos"]
    for field_id, coord in campos_map.items():
        if field_id.startswith("marca_"):
            # se manejan en _write_marks
            continue
        if field_id in fields:
            _set(ws, coord, fields[field_id])

    # Los grupos de marcas viven en `marcas:` del YAML, no aquí: así, cuando el
    # template vigente traiga la casilla que a este le falta (p. ej. el cuarto
    # segmento de productor), se añade una línea de config sin tocar código.
    for field_id, options in (cell_map.get("marcas") or {}).items():
        _write_marks(ws, fields, campos_map, field_id, options)


def _write_marks(ws, fields: dict, campos_map: dict, field_id: str, options: dict) -> None:
    """Marca con 'X' la opción seleccionada de un grupo (radio)."""
    value = (fields.get(field_id) or "").strip().lower()
    if not value:
        return
    target = options.get(value)
    if target and target in campos_map:
        _set(ws, campos_map[target], "X")


def _write_text_blocks(ws, fields: dict, cell_map: dict) -> None:
    if "forma_de_llegar" in fields:
        coord = cell_map["forma_de_llegar"]
        _set(ws, coord, fields["forma_de_llegar"])
        _wrap(ws, coord)
    if "justificacion_tecnica" in fields:
        coord = cell_map["justificacion_tecnica"]
        _set(ws, coord, fields["justificacion_tecnica"])
        _wrap(ws, coord)


def _wrap(ws, coord: str) -> None:
    cell = ws[coord]
    for mr in ws.merged_cells.ranges:
        if coord in mr:
            cell = ws.cell(row=mr.min_row, column=mr.min_col)
            break
    align = copy(cell.alignment) if cell.alignment else Alignment()
    align.wrap_text = True
    align.vertical = "top"
    cell.alignment = align


def _write_section_4(ws, fields: dict, cell_map: dict) -> None:
    for k in ("modalidad_pago_capital", "modalidad_pago_intereses"):
        if k in fields:
            _set(ws, cell_map[k], fields[k])


def _write_actividades_table(ws, fields: dict, tabla_map: dict) -> None:
    """Hasta 4 filas (49-52). Total en fila 53 ya tiene fórmula =SUM."""
    actividades = fields.get("actividades", []) or []
    cols = tabla_map["columnas"]
    first = tabla_map["primera_fila"]
    last = tabla_map["ultima_fila"]
    max_filas = last - first + 1

    if len(actividades) > max_filas:
        logger.warning("excel.actividades_truncadas", n=len(actividades), max=max_filas)
        actividades = actividades[:max_filas]

    for idx, act in enumerate(actividades):
        row = first + idx
        for key, col_letter in cols.items():
            value = act.get(key)
            if value is not None:
                _set(ws, f"{col_letter}{row}", value)


def _write_cronograma(ws, fields: dict, cell_map: dict) -> None:
    for k in ("cronograma_fecha_inicial", "cronograma_fecha_final"):
        if k in fields:
            _set(ws, cell_map[k], fields[k])


def _write_section_6_7(ws, fields: dict, cell_map: dict) -> None:
    # actividad_economica_codigo ahora se maneja en _write_digit_fields.
    for k in ("actividad_economica_descripcion", "fag_cobertura_pct"):
        if k in fields:
            _set(ws, cell_map[k], fields[k])

    fag = (fields.get("garantia_fag") or "").strip().lower()
    if fag in ("si", "sí", "true", "yes"):
        _set(ws, cell_map["marca_fag_si"], "X")
    elif fag in ("no", "false"):
        _set(ws, cell_map["marca_fag_no"], "X")


def _write_id_beneficiario(ws, fields: dict, id_map: dict) -> None:
    """Marca con X la fila correspondiente al tipo de ID y escribe el número.

    Sólo se rellena la fila del tipo elegido; las otras dos filas quedan en blanco.
    """
    if not id_map:
        return
    tipo = (fields.get("beneficiario_id_tipo") or "").strip().upper().replace(".", "")
    numero = fields.get("beneficiario_id_numero")
    aliases = {"CC": "cc", "CEDULA": "cc", "C C": "cc",
               "NIT": "nit",
               "CE": "ce", "CEDULA DE EXTRANJERIA": "ce", "C E": "ce"}
    key = aliases.get(tipo)
    if not key or key not in id_map:
        if tipo:
            logger.warning("excel.id_tipo_desconocido", tipo=tipo)
        return
    slot = id_map[key]
    _set(ws, slot["marca"], "X")
    if numero is not None:
        _set(ws, slot["numero"], numero)


def _write_digit_fields(ws, fields: dict, digit_map: dict) -> None:
    """Distribuye un código en celdas individuales (un dígito por celda).

    Si el valor es más corto que la cantidad de celdas, se padea con ceros
    a la izquierda (alinear=derecha) o a la derecha. Si es más largo, se trunca.
    """
    for field_id, spec in digit_map.items():
        value = fields.get(field_id)
        if value is None or value == "":
            continue
        celdas = spec.get("celdas", [])
        longitud = spec.get("longitud", len(celdas))
        alinear = spec.get("alinear", "derecha")
        s = str(value)
        if alinear == "derecha":
            s = s.rjust(longitud, "0")
        else:
            s = s.ljust(longitud, "0")
        s = s[:len(celdas)]
        for coord, digit in zip(celdas, s, strict=False):
            _set(ws, coord, digit)
