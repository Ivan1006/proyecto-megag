"""Parser de Excel: texto concatenado por hoja + tablas crudas.

Soporta `.xlsx`/`.xlsm` (openpyxl) y `.xls` antiguo BIFF (xlrd).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


def parse(path: Path) -> tuple[str, list[list[list[Any]]]]:
    ext = path.suffix.lower()
    if ext == ".xls":
        return _parse_xls(path)
    return _parse_xlsx(path)


def _parse_xlsx(path: Path) -> tuple[str, list[list[list[Any]]]]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    text_parts: list[str] = []
    tables: list[list[list[Any]]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        tables.append(rows)
        text_parts.append(f"[Hoja: {sheet_name}]")
        for row in rows:
            text_parts.append("\t".join("" if c is None else str(c) for c in row))
    return "\n".join(text_parts), tables


def _parse_xls(path: Path) -> tuple[str, list[list[list[Any]]]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    text_parts: list[str] = []
    tables: list[list[list[Any]]] = []
    for sheet in book.sheets():
        rows: list[list[Any]] = [
            [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        if not rows:
            continue
        tables.append(rows)
        text_parts.append(f"[Hoja: {sheet.name}]")
        for row in rows:
            text_parts.append("\t".join("" if c == "" else str(c) for c in row))
    return "\n".join(text_parts), tables
